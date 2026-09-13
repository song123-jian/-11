from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook


OUTPUT = Path("output/validation/office-samples")


def write_docx() -> None:
    document = Document()
    document.add_heading("文档办公工具验收样本", level=1)
    document.add_paragraph("这是 DOCX 转 PDF 的真实转换样本，包含中文标题、正文和表格。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "状态"
    table.cell(1, 1).text = "可转换"
    document.save(OUTPUT / "office-sample.docx")


def write_xlsx() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "验收数据"
    sheet.append(["项目", "数量", "说明"])
    sheet.append(["PDF 转换", 3, "DOCX / XLSX / PPTX"])
    sheet.append(["中文文本", 1000, "本地处理"])
    workbook.save(OUTPUT / "office-sample.xlsx")


def write_pptx() -> None:
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  <Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
</Types>"""
    root_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>"""
    presentation = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst>
  <p:sldSz cx="12192000" cy="6858000" type="screen16x9"/>
  <p:notesSz cx="6858000" cy="9144000"/>
</p:presentation>"""
    presentation_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
</Relationships>"""
    slide = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld><p:spTree>
    <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
    <p:grpSpPr/>
    <p:sp><p:nvSpPr><p:cNvPr id="2" name="标题"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="914400" y="914400"/><a:ext cx="10972800" cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="zh-CN" sz="3200"/><a:t>文档办公工具验收样本</a:t></a:r><a:endParaRPr lang="zh-CN"/></a:p></p:txBody></p:sp>
    <p:sp><p:nvSpPr><p:cNvPr id="3" name="正文"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="914400" y="2286000"/><a:ext cx="10972800" cy="2286000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="zh-CN" sz="2200"/><a:t>这是 PPTX 转 PDF 的真实转换样本，内容由本地生成。</a:t></a:r><a:endParaRPr lang="zh-CN"/></a:p></p:txBody></p:sp>
  </p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>"""
    target = OUTPUT / "office-sample.pptx"
    with ZipFile(target, "w", ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", content_types)
        package.writestr("_rels/.rels", root_rels)
        package.writestr("ppt/presentation.xml", presentation)
        package.writestr("ppt/_rels/presentation.xml.rels", presentation_rels)
        package.writestr("ppt/slides/slide1.xml", slide)


OUTPUT.mkdir(parents=True, exist_ok=True)
write_docx()
write_xlsx()
write_pptx()
print("generated", ", ".join(path.name for path in sorted(OUTPUT.glob("office-sample.*"))))
